"""文档加载：支持 PDF / DOCX / TXT / MD / CSV / XLSX。"""
import csv
import io
from typing import List

from langchain_core.documents import Document as LCDocument


def load_document_bytes(filename: str, data: bytes) -> List[LCDocument]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return _load_pdf(data)
    if ext == "docx":
        return _load_docx(data)
    if ext in ("txt", "md", "text", "markdown"):
        return _load_text(data)
    if ext == "csv":
        return _load_csv(data)
    if ext == "xlsx":
        return _load_xlsx(data)
    raise ValueError(f"不支持的文件类型: {ext or '未知'}")


def _load_pdf(data: bytes) -> List[LCDocument]:
    import fitz  # PyMuPDF

    docs: List[LCDocument] = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for i, page in enumerate(pdf):
            text = page.get_text()
            if text.strip():
                docs.append(LCDocument(page_content=text, metadata={"page": i + 1}))
    return docs


def _load_docx(data: bytes) -> List[LCDocument]:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    return [LCDocument(page_content=text, metadata={})] if text.strip() else []


def _load_text(data: bytes) -> List[LCDocument]:
    text = data.decode("utf-8", errors="ignore")
    return [LCDocument(page_content=text, metadata={})] if text.strip() else []


def _load_csv(data: bytes) -> List[LCDocument]:
    text = data.decode("utf-8-sig", errors="ignore")
    rows = list(csv.reader(io.StringIO(text)))
    lines = [" ".join(r).strip() for r in rows if any(cell.strip() for cell in r)]
    return [LCDocument(page_content="\n".join(lines), metadata={})] if lines else []


def _load_xlsx(data: bytes) -> List[LCDocument]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                lines.append(" ".join(str(c) for c in row if c is not None).strip())
    return [LCDocument(page_content="\n".join(lines), metadata={})] if lines else []
